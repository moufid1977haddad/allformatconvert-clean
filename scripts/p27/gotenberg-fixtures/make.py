"""P27 phase 3: two extra Word documents for the Gotenberg 8.36 / 8.37 comparison.
  math.docx          -- Office Math (OMML) equations in body text: 8.37's notes say DOCX math formulas were restored
                        (libreoffice-math re-added), so 8.36 may drop them.
  multilingual.docx  -- Greek, Arabic (right to left), Chinese, French accents, ligatures, in fonts Office documents
                        name (Calibri, Cambria, Arial, Microsoft YaHei): font substitution and shaping.
    python scripts/p27/gotenberg-fixtures/make.py
"""
import os

import docx
from docx.oxml import parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt

here = os.path.dirname(os.path.abspath(__file__))
M = 'xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"'


def omml(inner):
    return parse_xml(f'<m:oMathPara {M}><m:oMath>{inner}</m:oMath></m:oMathPara>')


def r(t):
    return f'<m:r><m:t>{t}</m:t></m:r>'


d = docx.Document()
d.add_heading('Equations (Office Math)', 1)
d.add_paragraph('The quadratic formula:')
p = d.add_paragraph()
p._p.append(omml(r('x=') + f'<m:f><m:num>{r("-b±")}<m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>'
                 + f'<m:sSup><m:e>{r("b")}</m:e><m:sup>{r("2")}</m:sup></m:sSup>{r("-4ac")}</m:e></m:rad></m:num>'
                 + f'<m:den>{r("2a")}</m:den></m:f>'))
d.add_paragraph('A sum and an integral:')
p = d.add_paragraph()
p._p.append(omml(f'<m:nary><m:naryPr><m:chr m:val="∑"/></m:naryPr><m:sub>{r("k=1")}</m:sub><m:sup>{r("n")}</m:sup>'
                 + f'<m:e>{r("k")}</m:e></m:nary>{r("=")}<m:f><m:num>{r("n(n+1)")}</m:num><m:den>{r("2")}</m:den></m:f>'))
p = d.add_paragraph()
p._p.append(omml(f'<m:nary><m:naryPr><m:chr m:val="∫"/></m:naryPr><m:sub>{r("0")}</m:sub><m:sup>{r("∞")}</m:sup>'
                 + f'<m:e><m:sSup><m:e>{r("e")}</m:e><m:sup>{r("-x²")}</m:sup></m:sSup>{r("dx")}</m:e></m:nary>'
                 + f'{r("=")}<m:f><m:num><m:rad><m:radPr><m:degHide m:val="1"/></m:radPr><m:deg/><m:e>{r("π")}</m:e></m:rad></m:num><m:den>{r("2")}</m:den></m:f>'))
d.add_paragraph('Text after the equations.')
d.save(os.path.join(here, 'math.docx'))

TEXTS = [
    ('Arial', None, 'Ελληνικά: η γλώσσα της Ελλάδας και της Κύπρου. Καλημέρα κόσμε.'),
    ('Arial', 'rtl', 'اللغة العربية لغة رسمية في اثنتين وعشرين دولة. مرحبا بالعالم.'),
    ('Microsoft YaHei', None, '中文：长期保存电子文档的国际标准。你好，世界！繁體中文：檔案長期保存。'),
    ('Cambria', None, 'Données numérisées : fidélité, intégrité, sécurité — à conserver « tel quel ». Œuvre, cœur, Noël.'),
    ('Calibri', None, 'office, affiche, efficient, fluffy, waffle, final, flow, section, action, attention, station.'),
]
d = docx.Document()
d.add_heading('Scripts and fonts', 1)
for font, mode, text in TEXTS:
    para = d.add_paragraph()
    run = para.add_run(text)
    run.font.name = font
    run.font.size = Pt(14)
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for k in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(k), font)
    if mode == 'rtl':
        rpr.append(rpr.makeelement(qn('w:rtl'), {}))
        ppr = para._p.get_or_add_pPr()
        ppr.append(ppr.makeelement(qn('w:bidi'), {}))
d.save(os.path.join(here, 'multilingual.docx'))
print('ok')
